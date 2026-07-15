from pptx.enum.shapes import MSO_SHAPE_TYPE


def update_table(shape, df):

    if not shape.has_table:
        return

    table = shape.table

    rows_required = len(df) + 1
    cols_required = len(df.columns)

    # Populate headers

    for col_index, column_name in enumerate(df.columns):

        if col_index < len(table.columns):

            table.cell(
                0,
                col_index
            ).text = str(column_name)

    # Populate values

    max_rows = min(
        len(table.rows) - 1,
        len(df)
    )

    for row_index in range(max_rows):

        for col_index in range(cols_required):

            if col_index < len(table.columns):

                value = df.iloc[
                    row_index,
                    col_index
                ]

                table.cell(
                    row_index + 1,
                    col_index
                ).text = str(value)